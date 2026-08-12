"""Results export. Runs against a read replica in production so a large report
cannot slow down a live exam."""

from __future__ import annotations

import io
import uuid
from datetime import UTC

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Answer, Exam, ExamAttempt, Question, QuestionType, Section, Student

HEADER_FONT = Font(bold=True)


async def build_results_workbook(db: AsyncSession, exam_id: uuid.UUID) -> tuple[str, bytes]:
    exam = await db.get(Exam, exam_id)
    if exam is None:
        raise ValueError("Exam not found")

    rows = await db.execute(
        select(ExamAttempt, Student)
        .join(Student, ExamAttempt.student_id == Student.id)
        .where(ExamAttempt.exam_id == exam_id)
        .order_by(ExamAttempt.total_score.desc().nullslast())
    )
    attempts = rows.all()

    counts = await db.execute(
        select(
            Answer.attempt_id,
            Question.type,
            Answer.is_correct,
        )
        .join(Question, Answer.question_id == Question.id)
        .join(ExamAttempt, Answer.attempt_id == ExamAttempt.id)
        .where(ExamAttempt.exam_id == exam_id)
    )
    per_attempt: dict[uuid.UUID, dict[str, int]] = {}
    for attempt_id, qtype, is_correct in counts.all():
        bucket = per_attempt.setdefault(
            attempt_id, {"attempted": 0, "correct": 0, "wrong": 0, "coding": 0}
        )
        bucket["attempted"] += 1
        if qtype is QuestionType.coding:
            bucket["coding"] += 1
        if is_correct is True:
            bucket["correct"] += 1
        elif is_correct is False:
            bucket["wrong"] += 1

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    headers = [
        "Rank",
        "Student ID",
        "Name",
        "Email",
        "Cohort",
        "Status",
        "Score",
        "Max Score",
        "Percentage",
        "Attempted",
        "Correct",
        "Wrong",
        "Coding Answers",
        "Started At",
        "Submitted At",
        "Duration (min)",
        "Tab Switches",
        "Grading Complete",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT

    for rank, (attempt, student) in enumerate(attempts, start=1):
        stats = per_attempt.get(attempt.id, {"attempted": 0, "correct": 0, "wrong": 0, "coding": 0})
        score = attempt.total_score
        max_score = attempt.max_score
        pct = round(score / max_score * 100, 2) if score is not None and max_score else None
        duration = None
        if attempt.submitted_at and attempt.started_at:
            duration = round(
                (
                    attempt.submitted_at.replace(tzinfo=attempt.submitted_at.tzinfo or UTC)
                    - attempt.started_at.replace(tzinfo=attempt.started_at.tzinfo or UTC)
                ).total_seconds()
                / 60,
                1,
            )
        ws.append(
            [
                rank,
                student.student_id,
                student.name,
                student.email,
                student.cohort,
                attempt.status.value,
                score,
                max_score,
                pct,
                stats["attempted"],
                stats["correct"],
                stats["wrong"],
                stats["coding"],
                attempt.started_at.replace(tzinfo=None) if attempt.started_at else None,
                attempt.submitted_at.replace(tzinfo=None) if attempt.submitted_at else None,
                duration,
                attempt.focus_loss_count,
                "Yes" if attempt.grading_complete else "No",
            ]
        )

    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(12, len(header) + 3)
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    safe_title = "".join(c if c.isalnum() or c in "-_ " else "_" for c in exam.title).strip()
    return f"{safe_title or 'exam'}_results.xlsx", buffer.getvalue()


async def build_attempts_workbook(db: AsyncSession, conditions: list, order_expr) -> tuple[str, bytes]:
    """Cross-exam export backing the Student Details filter bar — mirrors whatever
    filters produced the on-screen result set, with no pagination limit."""
    query = (
        select(ExamAttempt, Student, Exam)
        .join(Student, ExamAttempt.student_id == Student.id)
        .join(Exam, ExamAttempt.exam_id == Exam.id)
    )
    if conditions:
        query = query.where(and_(*conditions))

    rows = await db.execute(query.order_by(order_expr.nullslast()))
    results = rows.all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Attempts"
    headers = [
        "Student ID",
        "Name",
        "Email",
        "Cohort",
        "Exam",
        "Status",
        "Score",
        "Max Score",
        "Percentage",
        "Started At",
        "Submitted At",
        "Duration (min)",
        "Tab Switches",
        "Grading Complete",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT

    for attempt, student, exam in results:
        score = attempt.total_score
        max_score = attempt.max_score
        pct = round(score / max_score * 100, 2) if score is not None and max_score else None
        duration = None
        if attempt.submitted_at and attempt.started_at:
            duration = round(
                (
                    attempt.submitted_at.replace(tzinfo=attempt.submitted_at.tzinfo or UTC)
                    - attempt.started_at.replace(tzinfo=attempt.started_at.tzinfo or UTC)
                ).total_seconds()
                / 60,
                1,
            )
        ws.append(
            [
                student.student_id,
                student.name,
                student.email,
                student.cohort,
                exam.title,
                attempt.status.value,
                score,
                max_score,
                pct,
                attempt.started_at.replace(tzinfo=None) if attempt.started_at else None,
                attempt.submitted_at.replace(tzinfo=None) if attempt.submitted_at else None,
                duration,
                attempt.focus_loss_count,
                "Yes" if attempt.grading_complete else "No",
            ]
        )

    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(12, len(header) + 3)
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return "attempts_export.xlsx", buffer.getvalue()
